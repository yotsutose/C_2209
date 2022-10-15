import collections.abc # インポートしないとエラーが発生する
#from pptx import Presentation # プレゼンテーションを作成
from pptx.util import Inches  # インチ
#from pptx.enum.text import PP_ALIGN  # 中央揃えにする用
#from pptx.util import Cm, Pt # センチ、ポイント
import pptx
#import os

#cd = os.getcwd()
prs = pptx.Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[6])

for i in range(1,5):
    width = 1
    pic_left = Inches(width)
    pic_top = Inches(2)

    pic_height = Inches(3)

    image_path = "./input/IMG_000.png"
    slide.shapes.add_picture(image_path, pic_left, pic_top, height=pic_height)


prs.save("./test.pptx")